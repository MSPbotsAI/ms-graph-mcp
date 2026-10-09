"""Mail tools (Microsoft Graph message API).

Scope note: these tools cover one business flow — "what landed in this
mailbox, what is it about, and does it need a reply": list/search messages,
open one in full, read what an attachment says, send one. Moving, flagging,
deleting, replying in place and folder management are deliberately not
exposed — the read side exists to let an agent triage a mailbox and draft
follow-ups, not to act as a mail client.

Attachments are read, not downloaded: the bytes are fetched and parsed here
(spreadsheet -> rows, text -> string) and only the parsed content goes back.
Handing the raw file back as base64 would not fit the ~20,000-char result
cap for any real spreadsheet, and unlike drive items Graph has no
pre-authenticated download link for a message attachment to return instead.

Delegated vs app-only, and why it matters more here than anywhere else in
this server: delegated Mail.Read authorizes the SIGNED-IN user's own mailbox
and nothing else. Reading somebody else's mailbox under a delegated token
needs Mail.Read.Shared AND that mailbox actually delegated/shared to the
signed-in user in Exchange — the scope alone is not enough. Tenant-wide
access to every mailbox is not reachable delegated at all; that is what
ms-graph-app's application Mail.Read is for, and it should normally be
narrowed to specific mailboxes with an Exchange application access policy.
So under the delegated grant, omit user_id and you read your own mail; pass
someone else's and expect `unauthorized` unless the sharing is in place.
"""

import base64
import binascii
import csv
import datetime as dt
import io
import json
import re
import zipfile
from collections.abc import Callable
from typing import Annotated, Literal

from mcp.server.fastmcp import FastMCP
from mcp.types import ToolAnnotations
from openpyxl import load_workbook
from pydantic import Field

from .._json import MAX_CHARS, dump_json_capped, error_envelope
from ..api_client import GraphClient, GraphError
from ._common import NO_TOKEN, odata_quote

# Messages carry a bodyPreview each, so a page costs far more characters than
# a directory listing; 100 is the ceiling and the 20,000-char cap in _json.py
# still trims below it on chatty mailboxes.
_MAX_MESSAGES = 100
# A single message body can be hundreds of KB of quoted HTML thread history.
# dump_json_capped cannot rescue that on its own — a one-message dict has no
# list field to trim, so it would fall back to discarding the whole result and
# returning a "too large" notice. Truncating the body here keeps the useful
# head of the message instead.
_MAX_BODY_CHARS = 15_000

# Metadata only. Never bodyPreview's big sibling `body` — that is what
# graph_get_message is for.
_MESSAGE_FIELDS = (
    "id,conversationId,subject,from,toRecipients,ccRecipients,receivedDateTime,"
    "isRead,isDraft,hasAttachments,importance,flag,webLink,bodyPreview"
)

# Attachment reading. Exchange caps a message at ~150 MB, but nothing that big
# can be usefully summarized through a 20,000-char result anyway; refuse early
# instead of pulling it into memory.
_MAX_ATTACHMENT_READ_BYTES = 20 * 1024 * 1024
_MAX_SHEET_ROWS = 500
_MAX_CELL_CHARS = 300
_MAX_TEXT_CHARS = 15_000
_SPREADSHEET_EXTS = (".xlsx", ".xlsm")
_TEXT_EXTS = (".csv", ".tsv", ".txt", ".md", ".json", ".xml", ".html", ".htm", ".log")

_USER_ID_DESC = (
    "Mailbox to read (id (GUID) or userPrincipalName). Omit to use the token's own "
    "signed-in user — the normal case under the delegated grant, and the only mailbox "
    "it can read without extra sharing. An app-only token has no signed-in user, so it "
    "must always pass this."
)

# Inline attachments arrive as a base64 tool argument, so the calling agent has
# to hold (and emit) every byte in its own context first — same reasoning as
# sites.py's _MAX_WRITE_BYTES, and the per-file cap matches it. The combined
# cap is set by the gateway's front nginx (default client_max_body_size 1m):
# 700 KB decoded is ~933 KB of base64, leaving ~110 KB for the body, recipients
# and JSON-RPC envelope. Raise _MAX_TOTAL_ATTACHMENT_BYTES only after that
# nginx limit is raised. Graph's own ~4 MB sendMail request limit is far above.
_MAX_ATTACHMENT_BYTES = 500_000
_MAX_TOTAL_ATTACHMENT_BYTES = 700_000
_MAX_ATTACHMENTS = 5

_ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?(Z|[+-]\d{2}:\d{2})$")


def _build_attachments(attachments: list[dict]) -> list[dict] | str:
    """Validate inline attachments and map them to Graph fileAttachment dicts.

    Returns the list, or an invalid_argument envelope string — checked before
    any request goes out, so a bad attachment never sends a mail without it.
    """
    if len(attachments) > _MAX_ATTACHMENTS:
        return error_envelope(
            "invalid_argument",
            f"{_MAX_ATTACHMENTS} attachments max per call, got {len(attachments)}",
            False,
        )
    built: list[dict] = []
    total = 0
    for i, att in enumerate(attachments):
        if not isinstance(att, dict):
            return error_envelope("invalid_argument", f"attachments[{i}] must be an object", False)
        for key in ("filename", "content_base64"):
            if not att.get(key):
                return error_envelope(
                    "invalid_argument", f"attachments[{i}] missing required field '{key}'", False
                )
        filename = att["filename"]
        try:
            size = len(base64.b64decode(att["content_base64"], validate=True))
        except (binascii.Error, ValueError, TypeError) as e:
            return error_envelope(
                "invalid_argument",
                f"attachments[{i}] ('{filename}') is not valid base64: {e}",
                False,
            )
        if size > _MAX_ATTACHMENT_BYTES:
            return error_envelope(
                "invalid_argument",
                f"attachment '{filename}' is {size} bytes decoded, exceeds the "
                f"{_MAX_ATTACHMENT_BYTES:,}-byte (500 KB) per-file limit for inline attachments "
                "— compress the file or share a link instead",
                False,
            )
        total += size
        built.append(
            {
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": filename,
                "contentType": att.get("content_type") or "application/octet-stream",
                # Graph takes base64 as-is; no need to re-encode the decoded bytes.
                "contentBytes": att["content_base64"],
            }
        )
    if total > _MAX_TOTAL_ATTACHMENT_BYTES:
        return error_envelope(
            "invalid_argument",
            f"attachments total {total} bytes decoded, exceeds the "
            f"{_MAX_TOTAL_ATTACHMENT_BYTES:,}-byte (700 KB) combined limit; "
            "send in multiple emails or fewer/smaller attachments",
            False,
        )
    return built


def _cell(value) -> str | int | float | bool | None:
    """JSON-safe, length-capped cell value. Dates become ISO strings; long
    free-text cells are clipped so one notes column cannot eat the page."""
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    text = str(value)
    return text if len(text) <= _MAX_CELL_CHARS else text[:_MAX_CELL_CHARS] + "…"


def _trim_row(row) -> list:
    cells = [_cell(v) for v in row]
    while cells and cells[-1] in (None, ""):
        cells.pop()
    return cells


def _read_workbook(data: bytes, sheet: str | None, start_row: int, max_rows: int) -> dict | str:
    """Parse an .xlsx into one page of rows from one sheet.

    data_only=True returns the values Excel last calculated, not formula text —
    what a person sees in the cell. A workbook saved by a tool that never
    calculates (some exporters) has no cached values, so formula cells read None.
    """
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        return error_envelope("invalid_argument", f"Not a readable .xlsx workbook: {e}", False)
    try:
        names = wb.sheetnames
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in names:
            ws = wb[sheet]
        else:
            return error_envelope(
                "invalid_argument", f"No sheet named {sheet!r}; sheets are {names}", False
            )
        rows: list[list] = []
        total = 0
        # Blank rows are skipped rather than returned as [], and row numbers
        # count only non-blank rows, so start_row paging stays stable.
        for raw in ws.iter_rows(values_only=True):
            cells = _trim_row(raw)
            if not cells:
                continue
            total += 1
            if total >= start_row and len(rows) < max_rows:
                rows.append(cells)
        return {
            "sheets": names,
            "sheet": ws.title,
            "total_rows": total,
            "start_row": start_row,
            "returned_rows": len(rows),
            "has_more": start_row - 1 + len(rows) < total,
            "rows": rows,
        }
    finally:
        wb.close()


def _dump_rows(payload: dict) -> str:
    """dump_json_capped would trim `rows` too, but leave returned_rows/has_more
    describing the untrimmed page — the agent would then page past rows it never
    saw. Shrink the page here and keep the counters true."""
    rows = payload["rows"]
    lo, hi = 0, len(rows)
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if len(_page(payload, mid)) <= MAX_CHARS:
            lo = mid
        else:
            hi = mid - 1
    return _page(payload, lo)


def _page(payload: dict, n: int) -> str:
    page = dict(payload, rows=payload["rows"][:n], returned_rows=n)
    page["has_more"] = payload["start_row"] - 1 + n < payload["total_rows"]
    return json.dumps(page, separators=(",", ":"), ensure_ascii=False)


def _decode_text(data: bytes) -> str | None:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def _read_delimited(text: str, delimiter: str, start_row: int, max_rows: int) -> dict:
    rows: list[list] = []
    total = 0
    for raw in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = _trim_row(raw)
        if not cells:
            continue
        total += 1
        if total >= start_row and len(rows) < max_rows:
            rows.append(cells)
    return {
        "total_rows": total,
        "start_row": start_row,
        "returned_rows": len(rows),
        "has_more": start_row - 1 + len(rows) < total,
        "rows": rows,
    }


def register(mcp: FastMCP, client_factory: Callable[[], GraphClient | None]) -> None:

    @mcp.tool(
        annotations=ToolAnnotations(readOnlyHint=False, destructiveHint=False, idempotentHint=False)
    )
    async def graph_send_mail(
        to_recipients: Annotated[
            list[str],
            Field(description='List of To addresses, e.g. ["bob@contoso.com", "carol@contoso.com"].'),
        ],
        subject: Annotated[str, Field(description="Email subject line.")],
        body: Annotated[str, Field(description="Email body content.")],
        sender_id: Annotated[
            str | None,
            Field(
                description="Mailbox to send from (object ID or UPN). Omit to send as the token's own signed-in user — the normal case, since delegated Mail.Send only authorizes sending as that user. A different sender additionally needs Mail.Send.Shared and send-as rights."
            ),
        ] = None,
        body_content_type: Annotated[
            Literal["Text", "HTML"], Field(description='"Text" or "HTML".')
        ] = "Text",
        cc_recipients: Annotated[
            list[str] | None, Field(description="List of CC addresses.")
        ] = None,
        bcc_recipients: Annotated[
            list[str] | None, Field(description="List of BCC addresses.")
        ] = None,
        save_to_sent_items: Annotated[
            bool, Field(description="Whether to save a copy in Sent Items.")
        ] = True,
        attachments: Annotated[
            list[dict] | None,
            Field(
                description='Optional file attachments, up to 5: [{"filename": "quote.pdf", "content_base64": "<standard base64>", "content_type": "application/pdf"}]. content_type defaults to application/octet-stream. Max 500 KB per file and 700 KB total, decoded. If any attachment is invalid nothing is sent.'
            ),
        ] = None,
    ) -> str:
        """Send an email as an Entra ID user via Microsoft Graph.

        Irreversible once sent — no unsend. Confirm the exact recipients,
        subject, and body with the user before calling, especially when
        sender_id sends as someone other than the caller.
        """
        client = client_factory()
        if client is None:
            return NO_TOKEN

        def _addr_list(addresses: list[str]) -> list[dict]:
            return [{"emailAddress": {"address": addr}} for addr in addresses]

        message: dict = {
            "subject": subject,
            "body": {"contentType": body_content_type, "content": body},
            "toRecipients": _addr_list(to_recipients),
        }
        if cc_recipients:
            message["ccRecipients"] = _addr_list(cc_recipients)
        if bcc_recipients:
            message["bccRecipients"] = _addr_list(bcc_recipients)
        if attachments:
            built = _build_attachments(attachments)
            if isinstance(built, str):
                return built
            message["attachments"] = built

        payload = {
            "message": message,
            "saveToSentItems": save_to_sent_items,
        }

        # No sender_id -> /me — the only sender authorized under a delegated token.
        path = f"/users/{sender_id}/sendMail" if sender_id else "/me/sendMail"

        try:
            await client.post(path, payload)
            return dump_json_capped({"status": "sent"})
        except GraphError as e:
            return e.to_envelope()

    @mcp.tool(annotations=ToolAnnotations(readOnlyHint=True, idempotentHint=True))
    async def graph_list_messages(
        user_id: Annotated[str | None, Field(description=_USER_ID_DESC)] = None,
        folder: Annotated[
            str | None,
            Field(
                description='Mail folder to read: a well-known name ("inbox", "sentitems", '
                '"drafts", "archive", "junkemail", "deleteditems") or a folder id. Pass null to '
                "search every folder, which is what following a whole thread needs."
            ),
        ] = "inbox",
        search: Annotated[
            str | None,
            Field(
                description='Full-text search over the mailbox, e.g. "renewal quote" or '
                '"from:alice@contoso.com subject:invoice". Cannot be combined with the '
                "from_address / conversation_id / received_* / unread_only / has_attachments "
                "filters, and results come back by relevance rather than newest-first."
            ),
        ] = None,
        from_address: Annotated[
            str | None, Field(description="Keep only messages from this exact sender address.")
        ] = None,
        conversation_id: Annotated[
            str | None,
            Field(
                description="Keep only messages in one thread. Take it from a message's "
                "conversationId and pair it with folder=null to see the whole exchange, "
                "including replies that were sent rather than received."
            ),
        ] = None,
        received_after: Annotated[
            str | None,
            Field(description="Only messages received at or after this UTC time, e.g. 2026-09-01T00:00:00Z."),
        ] = None,
        received_before: Annotated[
            str | None,
            Field(description="Only messages received at or before this UTC time, e.g. 2026-09-08T00:00:00Z."),
        ] = None,
        unread_only: Annotated[
            bool, Field(description="Keep only unread messages.")
        ] = False,
        has_attachments: Annotated[
            bool | None, Field(description="Keep only messages with (True) or without (False) attachments.")
        ] = None,
        limit: Annotated[
            int,
            Field(description=f"Max messages to return (1-{_MAX_MESSAGES}).", ge=1, le=_MAX_MESSAGES),
        ] = 25,
    ) -> str:
        """List or search a mailbox, newest first, without message bodies.

        The triage entry point: returns sender, subject, received time,
        read state and a short bodyPreview per message, plus the id and
        conversationId needed to open one with graph_get_message or to pull
        the rest of its thread. One page only — narrow the filters rather
        than raising limit.
        """
        client = client_factory()
        if client is None:
            return NO_TOKEN

        filter_args = {
            "from_address": from_address,
            "conversation_id": conversation_id,
            "received_after": received_after,
            "received_before": received_before,
            "unread_only": unread_only or None,
            "has_attachments": has_attachments,
        }
        if search:
            conflicting = sorted(k for k, v in filter_args.items() if v is not None)
            if conflicting:
                return error_envelope(
                    "invalid_argument",
                    "Graph rejects $search combined with any other filter on messages. "
                    f"Drop {', '.join(conflicting)}, or fold the condition into the search "
                    'text itself (e.g. "from:alice@contoso.com renewal").',
                    False,
                )

        for name, value in (("received_after", received_after), ("received_before", received_before)):
            if value is not None and not _ISO_DATETIME.match(value):
                return error_envelope(
                    "invalid_argument",
                    f"{name} must be an ISO 8601 UTC datetime like 2026-09-01T00:00:00Z, got {value!r}.",
                    False,
                )

        base = f"/users/{user_id}" if user_id else "/me"
        path = f"{base}/mailFolders/{folder}/messages" if folder else f"{base}/messages"
        params: dict = {"$select": _MESSAGE_FIELDS, "$top": max(1, min(limit, _MAX_MESSAGES))}

        if search:
            # Graph wants the search term wrapped in double quotes inside the
            # query string; an embedded quote would close it early. $orderby is
            # not accepted alongside $search — results come back by relevance.
            params["$search"] = '"{}"'.format(search.replace('"', " ").strip())
        else:
            clauses: list[str] = []
            if from_address:
                clauses.append(f"from/emailAddress/address eq '{odata_quote(from_address)}'")
            if conversation_id:
                clauses.append(f"conversationId eq '{odata_quote(conversation_id)}'")
            if received_after:
                clauses.append(f"receivedDateTime ge {received_after}")
            if received_before:
                clauses.append(f"receivedDateTime le {received_before}")
            if unread_only:
                clauses.append("isRead eq false")
            if has_attachments is not None:
                clauses.append(f"hasAttachments eq {str(has_attachments).lower()}")
            if clauses:
                params["$filter"] = " and ".join(clauses)
            params["$orderby"] = "receivedDateTime desc"

        try:
            result = await client.get(path, params=params)
        except GraphError as e:
            return e.to_envelope()

        messages = result.get("value", []) if isinstance(result, dict) else []
        has_more = bool(isinstance(result, dict) and result.get("@odata.nextLink"))
        return dump_json_capped(
            {
                "mailbox": user_id or "me",
                "folder": folder or "all",
                "count": len(messages),
                "has_more": has_more,
                "messages": messages,
            }
        )

    @mcp.tool(annotations=ToolAnnotations(readOnlyHint=True, idempotentHint=True))
    async def graph_get_message(
        message_id: Annotated[
            str,
            Field(
                description="Message id from graph_list_messages — never guess it, list first."
            ),
        ],
        user_id: Annotated[str | None, Field(description=_USER_ID_DESC)] = None,
        body_format: Annotated[
            Literal["text", "html"],
            Field(
                description='"text" asks Exchange to flatten the body to plain text, which is '
                "several times cheaper to read and is almost always what you want; "
                '"html" keeps the original markup.'
            ),
        ] = "text",
        include_attachments: Annotated[
            bool,
            Field(
                description="Also list attachment ids, names, types and sizes (one extra "
                "call). To read a file's content pass its id to graph_read_mail_attachment."
            ),
        ] = False,
    ) -> str:
        """Read one message in full, including its body.

        Use after graph_list_messages has narrowed things down — a long
        thread's body is truncated here, and pulling bodies one at a time is
        what keeps a mailbox readable. The reply chain is usually quoted
        inside the body; sibling messages come from listing the
        conversationId instead.
        """
        client = client_factory()
        if client is None:
            return NO_TOKEN

        base = f"/users/{user_id}" if user_id else "/me"
        try:
            result = await client.get(
                f"{base}/messages/{message_id}",
                params={"$select": f"{_MESSAGE_FIELDS},body,replyTo"},
                # Exchange converts the stored body for us; asking for text here
                # is what avoids paying for a wall of HTML we would only strip.
                extra_headers={"Prefer": f'outlook.body-content-type="{body_format}"'},
            )
        except GraphError as e:
            return e.to_envelope()

        message = result if isinstance(result, dict) else {}
        body = message.get("body") or {}
        content = body.get("content") or ""
        if len(content) > _MAX_BODY_CHARS:
            message["body"] = {
                "contentType": body.get("contentType"),
                "content": content[:_MAX_BODY_CHARS],
                "truncated": True,
                "original_length": len(content),
            }

        payload: dict = {"mailbox": user_id or "me", "message": message}

        if include_attachments:
            try:
                # $select matters: without it Graph inlines contentBytes and a
                # single PDF turns the result into megabytes of base64.
                attached = await client.get(
                    f"{base}/messages/{message_id}/attachments",
                    params={"$select": "id,name,contentType,size,isInline"},
                )
                payload["attachments"] = (
                    attached.get("value", []) if isinstance(attached, dict) else []
                )
            except GraphError as e:
                return e.to_envelope()

        return dump_json_capped(payload)

    @mcp.tool(annotations=ToolAnnotations(readOnlyHint=True, idempotentHint=True))
    async def graph_read_mail_attachment(
        message_id: Annotated[str, Field(description="Message id from graph_list_messages.")],
        attachment_id: Annotated[
            str,
            Field(
                description="Attachment id from graph_get_message with include_attachments=true."
            ),
        ],
        user_id: Annotated[str | None, Field(description=_USER_ID_DESC)] = None,
        sheet: Annotated[
            str | None,
            Field(description="Excel only: sheet name to read. Omit for the first sheet."),
        ] = None,
        start_row: Annotated[
            int,
            Field(
                ge=1,
                description="Spreadsheet/CSV only: 1-based row to start from, counting "
                "non-blank rows. Use the previous call's start_row + returned_rows to page on.",
            ),
        ] = 1,
        max_rows: Annotated[
            int, Field(ge=1, le=_MAX_SHEET_ROWS, description="Spreadsheet/CSV rows per call.")
        ] = 200,
    ) -> str:
        """Read the content of a file attached to an email — Excel sheets as rows.

        .xlsx/.xlsm come back as rows of cell values from one sheet (first
        row is usually the header), with every sheet name listed;
        .csv/.tsv as rows; .txt/.json and similar as text. Large sheets
        page via start_row. Other types (.pdf/.docx/.xls) are rejected —
        the file itself is never returned.
        """
        client = client_factory()
        if client is None:
            return NO_TOKEN

        base = f"/users/{user_id}" if user_id else "/me"
        path = f"{base}/messages/{message_id}/attachments/{attachment_id}"
        try:
            # $select keeps contentBytes out: the size check must come first.
            meta = await client.get(path, params={"$select": "id,name,contentType,size"})
        except GraphError as e:
            return e.to_envelope()
        meta = meta if isinstance(meta, dict) else {}
        name = meta.get("name") or ""
        size = meta.get("size")
        kind = (meta.get("@odata.type") or "").rsplit(".", 1)[-1]
        info = {"name": name, "contentType": meta.get("contentType"), "size": size}

        if kind and kind != "fileAttachment":
            return error_envelope(
                "invalid_argument",
                f"{name!r} is an {kind} (an attached email/calendar item or a cloud link), "
                "not a file — nothing to read here.",
                False,
            )
        if isinstance(size, int) and size > _MAX_ATTACHMENT_READ_BYTES:
            return error_envelope(
                "invalid_argument",
                f"{name!r} is {size:,} bytes, over the "
                f"{_MAX_ATTACHMENT_READ_BYTES:,}-byte read limit.",
                False,
            )
        lower = name.lower()
        if not lower.endswith(_SPREADSHEET_EXTS + _TEXT_EXTS):
            return error_envelope(
                "invalid_argument",
                f"{name!r}: unsupported file type. Readable: "
                f"{', '.join(_SPREADSHEET_EXTS + _TEXT_EXTS)}. Legacy .xls is not supported.",
                False,
            )

        try:
            data = await client.get_content(f"{path}/$value")
        except GraphError as e:
            return e.to_envelope()

        if lower.endswith(_SPREADSHEET_EXTS):
            parsed = _read_workbook(data, sheet, start_row, max_rows)
            if isinstance(parsed, str):
                return parsed
            return _dump_rows({"attachment": info, **parsed})

        text = _decode_text(data)
        if text is None:
            return error_envelope(
                "invalid_argument", f"{name!r} is not valid UTF-8/UTF-16 text.", False
            )
        if lower.endswith((".csv", ".tsv")):
            delimiter = "\t" if lower.endswith(".tsv") else ","
            parsed = _read_delimited(text, delimiter, start_row, max_rows)
            return _dump_rows({"attachment": info, **parsed})
        content: dict = {"content": text[:_MAX_TEXT_CHARS]}
        if len(text) > _MAX_TEXT_CHARS:
            content.update(truncated=True, original_length=len(text))
        return dump_json_capped({"attachment": info, **content})
